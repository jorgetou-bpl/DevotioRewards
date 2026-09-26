# Customer insights routes — reads from the local customer_stats cache
# (backend/utils/boomerang.py, upsert_customer_stats), not Boomerangme
# directly. Kept separate from routes/operations.py (a different concern:
# aggregate/point-in-time stats derived from raw operations vs. a
# per-customer running cache) and from routes/customers.py (that one proxies
# Boomerangme's own /customers search, unscoped by workspace — a completely
# different data source and semantics from this file).

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from datetime import datetime, timezone
from typing import Optional
import io
import csv
import json
import logging
from utils.auth import get_current_user, require_workspace_admin
from utils.config import db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/customer-insights", tags=["customer-insights"])

AGE_BUCKETS = [
    ("18-24", 18, 25),
    ("25-34", 25, 35),
    ("35-44", 35, 45),
    ("45-54", 45, 55),
    ("55+", 55, 200),
]


@router.get("/age-distribution")
async def age_distribution(current_user: dict = Depends(require_workspace_admin)):
    """Bucket customer_stats.date_of_birth into age ranges. Computed in
    Python rather than a $dateDiff aggregation — the cache is one document
    per distinct customer (populated opportunistically on scan, no bulk
    backfill), small enough that a full pass is simpler and avoids a
    MongoDB-version dependency."""
    ws_id = current_user.get("workspace_id")
    total_customers = await db.customer_stats.count_documents({"workspace_id": ws_id})

    today = datetime.now(timezone.utc).date()
    buckets = {label: 0 for label, _, _ in AGE_BUCKETS}
    with_data = 0

    cursor = db.customer_stats.find(
        {"workspace_id": ws_id, "date_of_birth": {"$ne": None}},
        {"date_of_birth": 1}
    )
    async for doc in cursor:
        try:
            dob = datetime.fromisoformat(doc["date_of_birth"]).date()
        except (ValueError, TypeError):
            continue
        age = today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))
        with_data += 1
        for label, lo, hi in AGE_BUCKETS:
            if lo <= age < hi:
                buckets[label] += 1
                break

    return {
        "success": True,
        "buckets": [{"label": label, "count": buckets[label]} for label, _, _ in AGE_BUCKETS],
        "customers_with_data": with_data,
        "total_customers": total_customers
    }


VISIT_BUCKETS = [
    ("1 visita", 1, 2),
    ("2-3 visitas", 2, 4),
    ("4+ visitas", 4, 10**9),
]


@router.get("/visit-recurrence")
async def visit_recurrence(current_user: dict = Depends(get_current_user)):
    """Bucket customer_stats.total_visits into recurrence groups — a lifetime
    trait per customer (like age), not a per-period metric, so unlike
    /operations/summary this has no start_date/end_date. Open to every role
    (unlike age-distribution): visit counts aren't sensitive the way DOB is,
    and operators already see per-customer visit counts elsewhere."""
    ws_id = current_user.get("workspace_id")
    buckets = {label: 0 for label, _, _ in VISIT_BUCKETS}

    cursor = db.customer_stats.find({"workspace_id": ws_id}, {"total_visits": 1})
    async for doc in cursor:
        visits = doc.get("total_visits") or 0
        for label, lo, hi in VISIT_BUCKETS:
            if lo <= visits < hi:
                buckets[label] += 1
                break

    return {
        "success": True,
        "buckets": [{"label": label, "count": buckets[label]} for label, _, _ in VISIT_BUCKETS]
    }


MONTH_NAMES_ES = [
    "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"
]


@router.get("/new-customers-by-month")
async def new_customers_by_month(months: int = 6, current_user: dict = Depends(get_current_user)):
    """Cuenta de customer_stats.first_seen_at por mes, últimos N meses —
    mismo patrón que age-distribution/visit-recurrence (bucketing en Python,
    dataset chico). Va junto a Distribución por Edad en el pilar Visitas."""
    ws_id = current_user.get("workspace_id")
    months = max(3, min(months, 24))

    today = datetime.now(timezone.utc).date()
    month_keys = []
    y, m = today.year, today.month
    for _ in range(months):
        month_keys.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    month_keys.reverse()

    counts = {k: 0 for k in month_keys}
    earliest = month_keys[0]
    cursor = db.customer_stats.find(
        {"workspace_id": ws_id, "first_seen_at": {"$gte": earliest}},
        {"first_seen_at": 1}
    )
    async for doc in cursor:
        month_key = (doc.get("first_seen_at") or "")[:7]
        if month_key in counts:
            counts[month_key] += 1

    buckets = []
    for k in month_keys:
        year, month = k.split("-")
        buckets.append({"label": f"{MONTH_NAMES_ES[int(month) - 1]} {year}", "count": counts[k]})

    return {"success": True, "buckets": buckets}


ALLOWED_SORTS = {"last_seen_at", "total_visits", "first_seen_at", "customer_name", "total_purchase_sum"}

# Simple filters shared between the Customer Base list/export and (fast-
# follow) push notification targeting — a "segment" on either side is just
# a named preset built from these same field/operator/value triples, so the
# two features can never drift apart in what a given segment means.
FILTERABLE_FIELDS = {
    "total_visits", "total_purchase_sum", "stamps", "points_balance",
    "rewards_available", "last_seen_at", "first_seen_at"
}
OPERATOR_MAP = {"gte": "$gte", "lte": "$lte", "eq": "$eq"}


def _parse_filters(filters_json: Optional[str]) -> list:
    if not filters_json:
        return []
    try:
        parsed = json.loads(filters_json)
        return parsed if isinstance(parsed, list) else []
    except (ValueError, TypeError):
        return []


def _apply_filters(query: dict, filter_list: list) -> None:
    """Mutates `query` in place. "avg_spend" isn't a stored field (it's
    total_purchase_sum/total_visits), so it's the one case built as a Mongo
    $expr instead of a direct field comparison; everything else in
    FILTERABLE_FIELDS maps straight onto a customer_stats field."""
    expr_clauses = []
    for f in filter_list:
        field, operator, value = f.get("field"), f.get("operator"), f.get("value")
        mongo_op = OPERATOR_MAP.get(operator)
        if not mongo_op or value is None:
            continue
        if field == "avg_spend":
            expr_clauses.append({mongo_op: [
                {"$cond": [{"$gt": ["$total_visits", 0]}, {"$divide": ["$total_purchase_sum", "$total_visits"]}, 0]},
                value
            ]})
        elif field in FILTERABLE_FIELDS:
            query.setdefault(field, {})[mongo_op] = value
    if expr_clauses:
        query["$expr"] = {"$and": expr_clauses} if len(expr_clauses) > 1 else expr_clauses[0]


async def _phones_for_template(ws_id: str, template_id: str) -> list:
    """customer_stats has no template_id field (a customer's card_id is
    overwritten on every scan, not tracked per-template), so "customers of
    this card" is resolved from the operations they've actually logged
    against that template — same source Historial/Dashboard already filter
    by template_id."""
    return await db.operations.distinct(
        "customer_phone", {"workspace_id": ws_id, "template_id": str(template_id)}
    )


@router.get("/customers")
async def list_customers(
    page: int = 1,
    items_per_page: int = 25,
    sort_by: str = "last_seen_at",
    sort_dir: int = -1,
    template_id: Optional[str] = None,
    filters: Optional[str] = None,
    current_user: dict = Depends(require_workspace_admin)
):
    """Customer Base — the local customer_stats cache, not a live Boomerangme
    fetch, so this stays fast/cheap regardless of workspace size."""
    ws_id = current_user.get("workspace_id")
    if sort_by not in ALLOWED_SORTS:
        sort_by = "last_seen_at"
    skip = (page - 1) * items_per_page
    query = {"workspace_id": ws_id}
    if template_id:
        query["customer_phone"] = {"$in": await _phones_for_template(ws_id, template_id)}
    _apply_filters(query, _parse_filters(filters))

    total = await db.customer_stats.count_documents(query)
    cursor = db.customer_stats.find(query, {"_id": 0}) \
        .sort(sort_by, sort_dir).skip(skip).limit(items_per_page)
    customers = await cursor.to_list(length=items_per_page)

    return {
        "success": True,
        "customers": customers,
        "meta": {
            "total": total,
            "page": page,
            "items_per_page": items_per_page,
            "total_pages": (total + items_per_page - 1) // items_per_page
        }
    }


@router.get("/customers/export")
async def export_customers(
    format: str = "csv",
    sort_by: str = "last_seen_at",
    sort_dir: int = -1,
    filters: Optional[str] = None,
    current_user: dict = Depends(require_workspace_admin)
):
    """Export the full Customer Base (not just the current page) to CSV/XLSX.
    Registered before /customers/{phone} so "export" isn't matched as a
    phone number."""
    ws_id = current_user.get("workspace_id")
    if sort_by not in ALLOWED_SORTS:
        sort_by = "last_seen_at"

    query = {"workspace_id": ws_id}
    _apply_filters(query, _parse_filters(filters))
    cursor = db.customer_stats.find(query, {"_id": 0}).sort(sort_by, sort_dir)
    customers = await cursor.to_list(length=10000)

    headers = ["Nombre", "Teléfono", "Cliente desde", "Total Visitas", "Facturación Total", "Última Visita"]
    today_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    if format.lower() == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for c in customers:
            writer.writerow([
                c.get("customer_name", ""),
                c.get("customer_phone", ""),
                c.get("first_seen_at", ""),
                c.get("total_visits", 0),
                c.get("total_purchase_sum", 0),
                c.get("last_seen_at", "")
            ])
        output.seek(0)
        filename = f"clientes_{today_str}.csv"
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Content-Type": "text/csv; charset=utf-8"
            }
        )

    elif format.lower() == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "Clientes"

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="120627", end_color="120627", fill_type="solid")
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            for row_idx, c in enumerate(customers, 2):
                ws.cell(row=row_idx, column=1, value=c.get("customer_name", ""))
                ws.cell(row=row_idx, column=2, value=c.get("customer_phone", ""))
                ws.cell(row=row_idx, column=3, value=c.get("first_seen_at", ""))
                ws.cell(row=row_idx, column=4, value=c.get("total_visits", 0))
                ws.cell(row=row_idx, column=5, value=c.get("total_purchase_sum", 0))
                ws.cell(row=row_idx, column=6, value=c.get("last_seen_at", ""))

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                ws.column_dimensions[column].width = min(max_length + 2, 50)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            filename = f"clientes_{today_str}.xlsx"
            return StreamingResponse(
                iter([output.getvalue()]),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}"}
            )
        except ImportError:
            logger.warning("openpyxl not installed, falling back to CSV")
            raise HTTPException(status_code=400, detail="Formato XLSX no disponible. Use CSV.")

    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Use 'csv' o 'xlsx'.")


@router.get("/customers/{phone}")
async def get_customer_stats(phone: str, current_user: dict = Depends(require_workspace_admin)):
    ws_id = current_user.get("workspace_id")
    record = await db.customer_stats.find_one(
        {"workspace_id": ws_id, "customer_phone": phone}, {"_id": 0}
    )
    if not record:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    return {"success": True, "customer": record}
